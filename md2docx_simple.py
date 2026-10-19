#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Minimal Markdown -> .docx converter (headings, bold, bullets, pipe tables)."""
import os, re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_runs(par, text):
    """Handle **bold** and *italic* inline."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True
        else:
            par.add_run(tok)


def convert(md_path, docx_path):
    lines = open(md_path, encoding="utf-8").read().split("\n")
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.25

    i = 0
    n = len(lines)
    while i < n:
        ln = lines[i].rstrip()

        if not ln.strip():
            i += 1
            continue

        if ln.strip() == "---":
            i += 1
            continue

        # pipe table
        if ln.lstrip().startswith("|") and i + 1 < n and re.match(r"^\s*\|[\<PATH>", lines[i + 1]):
            header = [c.strip() for c in ln.strip().strip("|").split("|")]
            i += 2
            body = []
            while i < n and lines[i].lstrip().startswith("|"):
                body.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]
                cell.text = ""
                add_runs(cell.paragraphs[0], h)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in body:
                cells = t.add_row().cells
                for j in range(min(len(row), len(cells))):
                    cells[j].text = ""
                    add_runs(cells[j].paragraphs[0], row[j])
            doc.add_paragraph()
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", ln)
        if m:
            lvl = len(m.group(1))
            txt = m.group(2).strip()
            if lvl == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(txt); r.bold = True; r.font.size = Pt(15)
            else:
                p = doc.add_heading(level=min(lvl - 1, 4))
                add_runs(p, txt)
            i += 1
            continue

        if re.match(r"^\s*[-*]\s+", ln):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^\s*[-*]\s+", "", ln))
            i += 1
            continue

        if re.match(r"^\s*\d+\.\s+", ln):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\s*\d+\.\s+", "", ln))
            i += 1
            continue

        if ln.lstrip().startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, ln.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs(p, ln.strip())
        i += 1

    doc.save(docx_path)
    print("saved", docx_path)


if __name__ == "__main__":
    for md in sys.argv[1:]:
        out = os.path.splitext(md)[0] + ".docx"
        convert(md, out)
